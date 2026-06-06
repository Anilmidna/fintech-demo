"""
Run: python create_quiz.py
Creates: sigmapay_ahaslides_quiz.xlsx — ready to upload to AhaSlides
10 questions: 5 SQL (Module 1) + 5 Fraud Findings (Module 2)
SQL answers are actual queries, all under 100 chars.
"""
from openpyxl import load_workbook

wb = load_workbook("ahaslides-quiz-template_v2.xlsx")
ws = wb["Sheet1"]

for row in [9, 10, 11]:
    for col in range(1, 17):
        ws.cell(row=row, column=col).value = None

QUESTIONS = [

    # ── MODULE 1 — SQL (5 questions, answers are real queries) ─────

    {
        "num": 1,
        "question": "SQL Q1 — Count total rows in the transactions table. Which query is correct?",
        "answers": [
            "SELECT SUM(*) FROM transactions",
            "SELECT COUNT(*) FROM transactions",
            "SELECT ALL FROM transactions",
            "SELECT * FROM transactions",
        ],
        "correct": "2",
        "time": 30,
        "description": "COUNT(*) counts rows. SUM() adds numbers. Always COUNT for totals!",
    },
    {
        "num": 2,
        "question": "SQL Q2 — Show names of all merchants in Mumbai. Which query is correct?",
        "answers": [
            "SELECT name FROM merchants WHERE name='Mumbai'",
            "SELECT name FROM merchants WHERE city='Mumbai'",
            "SELECT city FROM merchants WHERE city='Mumbai'",
            "SELECT * FROM merchants ORDER BY city",
        ],
        "correct": "2",
        "time": 30,
        "description": "WHERE filters rows. You filter on city column, not name column. Easy mix-up!",
    },
    {
        "num": 3,
        "question": "SQL Q3 — Find total ₹ amount across ALL transactions. Which query is correct?",
        "answers": [
            "SELECT COUNT(amount) FROM transactions",
            "SELECT MAX(amount) FROM transactions",
            "SELECT SUM(amount) FROM transactions",
            "SELECT AVG(amount) FROM transactions",
        ],
        "correct": "3",
        "time": 30,
        "description": "SUM adds up all values. COUNT just counts rows. MAX finds the biggest. SUM is what you want here!",
    },
    {
        "num": 4,
        "question": "SQL Q4 — Show top 5 users by total spend. Which query is correct?",
        "answers": [
            "SELECT user_id, SUM(amount) FROM transactions GROUP BY user_id LIMIT 5",
            "SELECT user_id, SUM(amount) FROM transactions ORDER BY 2 DESC LIMIT 5",
            "SELECT user_id, SUM(amount) FROM transactions GROUP BY user_id ORDER BY 2 DESC LIMIT 5",
            "SELECT TOP 5 user_id, SUM(amount) FROM transactions",
        ],
        "correct": "3",
        "time": 30,
        "description": "Need GROUP BY to total per user, ORDER BY 2 DESC to rank highest first, LIMIT 5 to cap results.",
    },
    {
        "num": 5,
        "question": "SQL Q5 — Find all transactions that happened between midnight and 5 AM. Which query works?",
        "answers": [
            "SELECT * FROM transactions WHERE timestamp < '05:00'",
            "SELECT * FROM transactions WHERE strftime('%H',timestamp)<'05'",
            "SELECT * FROM transactions WHERE hour(timestamp)<5",
            "SELECT * FROM transactions WHERE timestamp LIKE '%-00-%'",
        ],
        "correct": "2",
        "time": 30,
        "description": "SQLite uses strftime() to extract hour from timestamp. '%H' gives the hour as 2-digit string.",
    },

    # ── MODULE 2 — FRAUD FINDINGS (5 questions) ────────────────────

    {
        "num": 6,
        "question": "Case 1 — Speed Merchant: How many transactions did U005 make in under 30 minutes?",
        "answers": ["3 transactions", "8 transactions", "18 transactions", "42 transactions"],
        "correct": "3",
        "time": 20,
        "description": "18 txns in 30 min! Normal users don't do that. Classic high-frequency fraud signal.",
    },
    {
        "num": 7,
        "question": "Case 2 — Round Numbers: Which amounts did U008 always send?",
        "answers": [
            "₹347, ₹892 — random amounts",
            "₹500, ₹1000, ₹2000 — always round",
            "₹9,999 — just below the limit",
            "₹1, ₹5 — tiny test amounts",
        ],
        "correct": "2",
        "time": 20,
        "description": "Round numbers = automated fraud script. Real humans buy groceries for ₹347, not exactly ₹500.",
    },
    {
        "num": 8,
        "question": "Case 3 — Threshold Dodger: Why does U015 always send ₹9,800–₹9,999, never ₹10,000+?",
        "answers": [
            "Their UPI app has a ₹10,000 cap",
            "They are a careful spender",
            "Above ₹10,000 triggers KYC — deliberately avoiding it",
            "Pure coincidence",
        ],
        "correct": "3",
        "time": 30,
        "description": "Called structuring or threshold skirting — a classic money laundering technique used worldwide.",
    },
    {
        "num": 9,
        "question": "Case 4 — Sleeping Giant: 'Dormant' users suddenly made 8+ txns/day AND their account_age_days was just 3–8. What's really going on?",
        "answers": [
            "Genuine old customers, just back after a break",
            "A system glitch — ignore both numbers",
            "Freshly created mule accounts disguised as dormant, now moving stolen money",
            "New customers who got lucky with a big offer",
        ],
        "correct": "3",
        "time": 30,
        "description": "Real dormant accounts are OLD accounts gone quiet. These are 3–8 days old — brand new. Label them 'dormant' to look harmless, then burst into activity to launder money. That's a mule account.",
    },
    {
        "num": 10,
        "question": "Case 5 — Card Tester: After U018 sent several tiny test amounts (all under ₹10) — what was their very next transaction?",
        "answers": [
            "₹500 — a normal purchase",
            "₹9,999 — just under KYC limit",
            "Account was blocked immediately",
            "₹48,000 — the big hit after confirming card works",
        ],
        "correct": "4",
        "time": 20,
        "description": "Tiny amounts = card test. ₹48,000 = the real theft. You just caught a fraudster using SQL!",
    },
]

# Write to sheet
for i, q in enumerate(QUESTIONS):
    row = 9 + i
    ws.cell(row=row, column=1).value = q["num"]
    ws.cell(row=row, column=2).value = q["question"]
    for j, ans in enumerate(q["answers"]):
        ws.cell(row=row, column=3 + j).value = ans
    ws.cell(row=row, column=11).value = q["time"]
    ws.cell(row=row, column=12).value = q["correct"]
    ws.cell(row=row, column=13).value = "Yes"
    ws.cell(row=row, column=14).value = 100
    ws.cell(row=row, column=15).value = 1000
    ws.cell(row=row, column=16).value = q["description"]

# Char limit check
print("── Character check ──────────────────────────")
ok = True
for q in QUESTIONS:
    if len(q["question"]) > 300:
        print(f"❌ Q{q['num']} question: {len(q['question'])} chars (max 300)")
        ok = False
    for j, ans in enumerate(q["answers"]):
        if len(ans) > 150:
            print(f"❌ Q{q['num']} answer {j+1}: {len(ans)} chars (max 150)")
            ok = False
if ok:
    print("✅ All questions and answers within AhaSlides limits")

output = "sigmapay_ahaslides_quiz.xlsx"
wb.save(output)
print(f"\n✅ Saved: {output}  ({len(QUESTIONS)} questions)")
print(f"   Q1–Q5  : SQL queries as answer options (30 sec each)")
print(f"   Q6–Q10 : Fraud case findings (20–30 sec each)")
print(f"   Upload : AhaSlides → Quiz → Import")
